#!/usr/bin/env python3
"""
build_deck.py
=============
Constructs the complete 12-slide professional academic PowerPoint (.pptx)
presentation for the M.Tech mid-semester panel evaluation.

Presenter: Anvesh Dhiraj Shivhare
Institution: National Forensic Sciences University (NFSU), Gandhinagar
Topic: AI-Assisted Detection of LOLBin-Based Privilege Escalation
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# ------------------------------------------------------------------------------
# PALETTE DEFINITIONS (Professional Academic Theme)
# ------------------------------------------------------------------------------
COLOR_NAVY_DARK   = RGBColor(15, 23, 42)     # #0F172A - Deep primary text & title
COLOR_NAVY_MID    = RGBColor(30, 58, 138)    # #1E3A8A - Primary academic blue
COLOR_ACCENT_BLUE = RGBColor(37, 99, 235)    # #2563EB - Highlights & key metrics
COLOR_SLATE_BODY  = RGBColor(51, 65, 85)     # #334155 - High readability body text
COLOR_SLATE_MUTED = RGBColor(100, 116, 139)  # #64748B - Footers, subtitles, notes
COLOR_BG_CARD     = RGBColor(255, 255, 255)  # #FFFFFF - White card background
COLOR_BG_PAGE     = RGBColor(248, 250, 252)  # #F8FAFC - Slide background
COLOR_BORDER_CARD = RGBColor(226, 232, 240)  # #E2E8F0 - Subtle card border

# BADGE STYLING
# Real Data (Emerald)
BADGE_REAL_BG     = RGBColor(220, 252, 231)  # #DCFCE7
BADGE_REAL_BORDER = RGBColor(22, 163, 74)    # #16A34A
BADGE_REAL_TEXT   = RGBColor(21, 128, 61)    # #15803D

# Synthetic Data (Amber)
BADGE_SYNTH_BG     = RGBColor(254, 243, 199) # #FEF3C7
BADGE_SYNTH_BORDER = RGBColor(217, 119, 6)   # #D97706
BADGE_SYNTH_TEXT   = RGBColor(180, 83, 9)    # #B45309

# Design / Foundation (Slate)
BADGE_NEUTRAL_BG     = RGBColor(241, 245, 249) # #F1F5F9
BADGE_NEUTRAL_BORDER = RGBColor(148, 163, 184) # #94A3B8
BADGE_NEUTRAL_TEXT   = RGBColor(71, 85, 105)   # #475569

# Honest Limitation (Red)
BADGE_LIMIT_BG     = RGBColor(254, 226, 226) # #FEE2E2
BADGE_LIMIT_BORDER = RGBColor(220, 38, 38)   # #DC2626
BADGE_LIMIT_TEXT   = RGBColor(185, 28, 28)   # #B91C1C

FONT_HEADING = "Helvetica"
FONT_BODY    = "Helvetica"


def create_deck(output_pptx="presentation/LOLBin_PrivEsc_Detection_MidSem.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_page_decorations(slide, slide_num, total_slides=12):
        """Adds consistent academic header, footer, and slide number."""
        # Top banner thin accent line
        accent_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08)
        )
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = COLOR_NAVY_MID
        accent_line.line.fill.background()

        # Footer divider line
        f_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.0), Inches(11.733), Inches(0.015)
        )
        f_line.fill.solid()
        f_line.fill.fore_color.rgb = COLOR_BORDER_CARD
        f_line.line.fill.background()

        # Footer text left
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(9.0), Inches(0.35))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = "AI-Assisted Detection of LOLBin-Based Privilege Escalation | NFSU M.Tech Cyber Security"
        p.font.name = FONT_BODY
        p.font.size = Pt(9)
        p.font.color.rgb = COLOR_SLATE_MUTED

        # Footer text right (slide number)
        tx_num = slide.shapes.add_textbox(Inches(10.5), Inches(7.05), Inches(2.0), Inches(0.35))
        tf_num = tx_num.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.text = f"Slide {slide_num} of {total_slides}"
        p_num.font.name = FONT_BODY
        p_num.font.size = Pt(9)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_SLATE_MUTED

    def add_header(slide, title, subtitle, badge_type, badge_text):
        """Standardized title, subtitle, and provenance badge."""
        # Badge
        badge_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(4.5), Inches(0.32)
        )
        badge_box.line.width = Pt(1)

        if badge_type == "real":
            badge_box.fill.solid()
            badge_box.fill.fore_color.rgb = BADGE_REAL_BG
            badge_box.line.color.rgb = BADGE_REAL_BORDER
            text_color = BADGE_REAL_TEXT
        elif badge_type == "synth":
            badge_box.fill.solid()
            badge_box.fill.fore_color.rgb = BADGE_SYNTH_BG
            badge_box.line.color.rgb = BADGE_SYNTH_BORDER
            text_color = BADGE_SYNTH_TEXT
        elif badge_type == "limit":
            badge_box.fill.solid()
            badge_box.fill.fore_color.rgb = BADGE_LIMIT_BG
            badge_box.line.color.rgb = BADGE_LIMIT_BORDER
            text_color = BADGE_LIMIT_TEXT
        else:
            badge_box.fill.solid()
            badge_box.fill.fore_color.rgb = BADGE_NEUTRAL_BG
            badge_box.line.color.rgb = BADGE_NEUTRAL_BORDER
            text_color = BADGE_NEUTRAL_TEXT

        tf_badge = badge_box.text_frame
        tf_badge.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_badge = tf_badge.paragraphs[0]
        p_badge.text = badge_text
        p_badge.font.name = FONT_BODY
        p_badge.font.size = Pt(8.5)
        p_badge.font.bold = True
        p_badge.font.color.rgb = text_color
        p_badge.alignment = PP_ALIGN.CENTER

        # Title
        tx_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.80), Inches(11.733), Inches(0.55))
        tf_title = tx_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY_DARK

        # Subtitle
        tx_sub = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.35))
        tf_sub = tx_sub.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = COLOR_SLATE_MUTED

    # ==========================================================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================================================
    slide1 = prs.slides.add_slide(blank_layout)

    # Top accent bar
    top_bar = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_NAVY_MID
    top_bar.line.fill.background()

    # Left decorative vertical accent bar
    stripe = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.12), Inches(5.1)
    )
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    stripe.line.fill.background()

    # Title text box
    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.15), Inches(11.2), Inches(2.3))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p_eval = tf1.paragraphs[0]
    p_eval.text = "M.TECH CYBER SECURITY — MID-SEMESTER DISSERTATION EVALUATION"
    p_eval.font.name = FONT_BODY
    p_eval.font.size = Pt(11)
    p_eval.font.bold = True
    p_eval.font.color.rgb = COLOR_ACCENT_BLUE
    p_eval.space_after = Pt(10)

    p_proj = tf1.add_paragraph()
    p_proj.text = "AI-Assisted Detection of LOLBin-Based\nPrivilege Escalation"
    p_proj.font.name = FONT_HEADING
    p_proj.font.size = Pt(32)
    p_proj.font.bold = True
    p_proj.font.color.rgb = COLOR_NAVY_DARK
    p_proj.space_after = Pt(12)

    p_phase = tf1.add_paragraph()
    p_phase.text = "Phase 1 of 2: Generic Behavioral Detection Foundation"
    p_phase.font.name = FONT_BODY
    p_phase.font.size = Pt(16)
    p_phase.font.color.rgb = COLOR_NAVY_MID
    p_phase.font.bold = True

    # Divider line
    div1 = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(3.85), Inches(10.5), Inches(0.015)
    )
    div1.fill.solid()
    div1.fill.fore_color.rgb = COLOR_BORDER_CARD
    div1.line.fill.background()

    # 4 Metadata Cards
    meta_cards = [
        ("PRESENTER", "Anvesh Dhiraj Shivhare\nM.Tech Cyber Security, Semester III"),
        ("INSTITUTION", "School of Cyber Security & Digital Forensics\nNational Forensic Sciences University, Gandhinagar"),
        ("RESEARCH TRACK", "Defensive EDR / Host Behavioral Anomaly Detection\nUnsupervised Learning + Zero-Shot Explainability"),
        ("SUBMISSION CONTEXT", "October 2026\nAcademic Year 2026–2027 (Mid-Semester Panel)")
    ]

    for i, (label, val) in enumerate(meta_cards):
        col = i % 2
        row = i // 2
        left = Inches(1.2 + col * 5.4)
        top = Inches(4.15 + row * 1.35)

        card = slide1.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.1), Inches(1.18)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(248, 250, 252)
        card.line.color.rgb = COLOR_BORDER_CARD
        card.line.width = Pt(1)

        c_tf = card.text_frame
        c_tf.vertical_anchor = MSO_ANCHOR.TOP
        c_tf.margin_left = Inches(0.2)
        c_tf.margin_top = Inches(0.15)

        p_l = c_tf.paragraphs[0]
        p_l.alignment = PP_ALIGN.LEFT
        p_l.text = label
        p_l.font.name = FONT_BODY
        p_l.font.size = Pt(8.5)
        p_l.font.bold = True
        p_l.font.color.rgb = COLOR_ACCENT_BLUE

        p_v = c_tf.add_paragraph()
        p_v.alignment = PP_ALIGN.LEFT
        p_v.text = val
        p_v.font.name = FONT_BODY
        p_v.font.size = Pt(10)
        p_v.font.color.rgb = COLOR_NAVY_DARK
        p_v.space_before = Pt(3)

    add_page_decorations(slide1, 1)

    # ==========================================================================
    # SLIDE 2: PROBLEM STATEMENT
    # ==========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(
        slide2,
        title="Problem Formulation: The Living-off-the-Land Threat",
        subtitle="Why legitimate dual-use administrative binaries evade signature-based AV and static EDR",
        badge_type="neutral",
        badge_text="◆ THREAT MODEL & STRATEGIC PIVOT"
    )

    p2_points = [
        (
            "Adversarial Abuse of Dual-Use Utilities (LOLBins / GTFOBins)",
            "Modern threat actors avoid compiling custom malware binaries. Instead, they weaponize pre-installed, digitally signed system utilities (e.g., powershell.exe, certutil.exe, find, awk, tar, sudo) for post-exploitation and privilege escalation (PrivEsc)."
        ),
        (
            "Failure of Static Signatures & Crippling Rule Fatigue",
            "Because LOLBins inherit valid OS signatures and run from trusted system paths, traditional signature-based AV is ineffective. Static EDR rules face an irreconcilable tradeoff: either trigger severe False Positive alert fatigue on routine administration or permit intrusions."
        ),
        (
            "The Defensive Research Gap (Our Strategic Pivot)",
            "Offensive AI automation (auto-exploitation, LLM payload mutation) is saturated in literature (PentestGPT, hackingBuddyGPT). Detecting subtle dual-use commands via unsupervised behavioral anomaly detection paired with zero-shot LLM explainability is an open, defensible gap."
        ),
        (
            "The Sequence Blind Spot (Isolated vs. Chained Execution)",
            "A dual-use utility (e.g., 'whoami' or 'id') looks completely benign in isolation. Malicious intent only manifests across consecutive command transitions within an attack session, requiring sequence-aware behavioral analysis rather than isolated inspection."
        )
    ]

    for i, (heading, desc) in enumerate(p2_points):
        top_pos = Inches(1.85 + i * 1.22)
        card = slide2.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.733), Inches(1.08)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_BORDER_CARD
        card.line.width = Pt(1)

        # Left indicator tag
        tag = slide2.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), top_pos, Inches(0.1), Inches(1.08)
        )
        tag.fill.solid()
        tag.fill.fore_color.rgb = COLOR_NAVY_MID
        tag.line.fill.background()

        ctf = card.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.12)

        p_h = ctf.paragraphs[0]
        p_h.alignment = PP_ALIGN.LEFT
        p_h.text = f"{i+1}. {heading}"
        p_h.font.name = FONT_HEADING
        p_h.font.size = Pt(11.5)
        p_h.font.bold = True
        p_h.font.color.rgb = COLOR_NAVY_DARK

        p_d = ctf.add_paragraph()
        p_d.alignment = PP_ALIGN.LEFT
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_SLATE_BODY
        p_d.space_before = Pt(3)

    add_page_decorations(slide2, 2)

    # ==========================================================================
    # SLIDE 3: RESEARCH FOUNDATION
    # ==========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(
        slide3,
        title="Theoretical & Literature Foundation",
        subtitle="Mapping peer-reviewed foundational literature directly to each architectural design decision",
        badge_type="neutral",
        badge_text="◆ LITERATURE TAXONOMY (IEEE / APA CITATIONS)"
    )

    citations = [
        (
            "Isolation Forest Algorithm",
            "Liu, Ting, Zhou (IEEE ICDM 2008)",
            "Exploits fundamental property: 'anomalies are few and different'. Isolates anomalies via random recursive partitioning without requiring labeled attack training traces."
        ),
        (
            "Why Unsupervised Learning",
            "Stamp (2022, arXiv:2208.12836)",
            "Demonstrates severe scarcity of labeled malicious LOLBin data. Proves supervised models overfit to seen syntax and collapse when facing novel post-exploitation techniques."
        ),
        (
            "Anomaly Detector + LLM Pairing",
            "Ali & Kostakos (2023, HuntGPT)",
            "Establishes architectural precedent combining statistical ML anomaly detection with zero-shot LLM explainability to translate mathematical deviations into actionable SOC triage."
        ),
        (
            "LLM Grounding & Hallucination Guard",
            "Houssel et al. (2024, arXiv:2408.04342)",
            "Defines methodology for strictly grounding LLM reasoning in empirical feature deviations and z-scores, prohibiting ungrounded external assumptions."
        ),
        (
            "Sliding-Window Command Context",
            "Springer (2026)",
            "Proves empirically that evaluating a sliding window of consecutive commands significantly outperforms isolated single-command inspection on Linux endpoint telemetry."
        ),
        (
            "Anomaly Weakness on Stealth / Obfuscation",
            "Systematic Literature Review (SLR 2026)",
            "Exhaustive review establishing that Isolation Forest excels on overt anomalies but exhibits known degradation on stealthy/obfuscated intrusions—grounding our Tier 3 test."
        )
    ]

    for i, (comp, paper, impact) in enumerate(citations):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.85 + row * 1.62)

        card = slide3.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(1.5)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_BORDER_CARD
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.margin_left = Inches(0.2)
        ctf.margin_top = Inches(0.12)

        p_c = ctf.paragraphs[0]
        p_c.alignment = PP_ALIGN.LEFT
        p_c.text = comp
        p_c.font.name = FONT_HEADING
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = COLOR_NAVY_MID

        p_p = ctf.add_paragraph()
        p_p.alignment = PP_ALIGN.LEFT
        p_p.text = f"Anchor: {paper}"
        p_p.font.name = FONT_BODY
        p_p.font.size = Pt(9.5)
        p_p.font.bold = True
        p_p.font.color.rgb = COLOR_ACCENT_BLUE
        p_p.space_before = Pt(2)

        p_i = ctf.add_paragraph()
        p_i.alignment = PP_ALIGN.LEFT
        p_i.text = impact
        p_i.font.name = FONT_BODY
        p_i.font.size = Pt(9)
        p_i.font.color.rgb = COLOR_SLATE_BODY
        p_i.space_before = Pt(3)

    add_page_decorations(slide3, 3)

    # ==========================================================================
    # SLIDE 4: SYSTEM ARCHITECTURE
    # ==========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(
        slide4,
        title="System Architecture: Behavioral Detection Pipeline",
        subtitle="End-to-end ingestion, feature representation, unsupervised scoring, and grounded LLM attribution",
        badge_type="neutral",
        badge_text="◆ END-TO-END PIPELINE ARCHITECTURE"
    )

    stages = [
        (
            "1. Telemetry Ingestion",
            "AuditdTelemetrySource",
            [
                "Ingests raw Linux auditd streams",
                "Correlates SYSCALL, EXECVE, PATH, CWD",
                "Decodes hex proctitles into plaintext",
                "Extensible adapter architecture"
            ],
            "REAL DATA (auditd)\n& Synthetic Adapter"
        ),
        (
            "2. Feature Engineering",
            "TelemetryFeatureExtractor",
            [
                "Lineage rarity (parent-child co-occurrence)",
                "Shannon entropy & special char ratio",
                "Sensitive path touch & privilege flags",
                "Sliding Command Window (N=3)"
            ],
            "Real Fitted Statistics\n(Zero Retraining)"
        ),
        (
            "3. Anomaly Detection",
            "Isolation Forest (Unsupervised)",
            [
                "150 Isolation Trees (contamination=0.025)",
                "Trained strictly on benign baseline",
                "Continuous standardized feature space",
                "Calibrated decision boundary = 0.50"
            ],
            "Real Auditd Baseline\n(25,495 Events)"
        ),
        (
            "4. SOC Explainability",
            "LLM Attribution & Triage",
            [
                "Zero-shot Tier-3 SOC analyst prompt",
                "Consumes JSON feature anomaly payload",
                "Outputs ATT&CK ID & confidence score",
                "Deterministic offline fallback engine"
            ],
            "Architecture Validated\n(No Label Leakage)"
        )
    ]

    for i, (title, comp, desc_bullets, status) in enumerate(stages):
        left = Inches(0.8 + i * 2.98)
        card = slide4.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.85), Inches(2.8), Inches(4.3)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_BORDER_CARD
        card.line.width = Pt(1)

        # Header tag
        tag = slide4.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, Inches(1.85), Inches(2.8), Inches(0.42)
        )
        tag.fill.solid()
        tag.fill.fore_color.rgb = COLOR_NAVY_MID
        tag.line.fill.background()
        tf_tag = tag.text_frame
        tf_tag.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_t = tf_tag.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(255, 255, 255)
        p_t.alignment = PP_ALIGN.CENTER

        # Body text box
        tx_body = slide4.shapes.add_textbox(left + Inches(0.12), Inches(2.32), Inches(2.56), Inches(2.7))
        ctf = tx_body.text_frame
        ctf.word_wrap = True

        p_comp = ctf.paragraphs[0]
        p_comp.text = comp
        p_comp.font.name = FONT_BODY
        p_comp.font.size = Pt(9.5)
        p_comp.font.bold = True
        p_comp.font.color.rgb = COLOR_ACCENT_BLUE

        for b in desc_bullets:
            p_b = ctf.add_paragraph()
            p_b.text = f"• {b}"
            p_b.font.name = FONT_BODY
            p_b.font.size = Pt(8.5)
            p_b.font.color.rgb = COLOR_SLATE_BODY
            p_b.space_before = Pt(4)

        # Status badge at bottom of card
        s_box = slide4.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.15), Inches(5.38), Inches(2.5), Inches(0.62)
        )
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
        s_box.line.color.rgb = COLOR_BORDER_CARD
        s_tf = s_box.text_frame
        s_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_s = s_tf.paragraphs[0]
        p_s.text = status
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(8)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_NAVY_DARK
        p_s.alignment = PP_ALIGN.CENTER

    # Arrow connectors between stages
    for i in range(3):
        arrow = slide4.shapes.add_textbox(Inches(3.55 + i * 2.98), Inches(3.7), Inches(0.3), Inches(0.3))
        tf_a = arrow.text_frame
        p_a = tf_a.paragraphs[0]
        p_a.text = "➔"
        p_a.font.size = Pt(16)
        p_a.font.bold = True
        p_a.font.color.rgb = COLOR_ACCENT_BLUE

    # Provenance clarification note at bottom
    note_box = slide4.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(11.733), Inches(0.6))
    tf_n = note_box.text_frame
    p_n = tf_n.paragraphs[0]
    p_n.text = "Data Provenance Distinction: Detection & Feature Pipeline is trained and evaluated on REAL Ubuntu auditd telemetry. Multi-tier adversarial benchmarking (Tiers 1–3) was validated against synthetic lab streams during architectural development."
    p_n.font.name = FONT_BODY
    p_n.font.size = Pt(9)
    p_n.font.color.rgb = COLOR_SLATE_MUTED

    add_page_decorations(slide4, 4)

    # ==========================================================================
    # SLIDE 5: IMPLEMENTATION STATUS
    # ==========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(
        slide5,
        title="Implementation Status: Phase 1 Scope & Boundary",
        subtitle="Transparent accounting of what is delivered and operational vs. synthetic and Phase 2 targets",
        badge_type="neutral",
        badge_text="◆ PROJECT MILESTONES & RIGOROUS SCOPING"
    )

    status_cols = [
        (
            "BUILT & OPERATIONAL ON REAL DATA",
            "Phase 1 Delivered (Verified on Linux Host)",
            BADGE_REAL_BG,
            BADGE_REAL_BORDER,
            BADGE_REAL_TEXT,
            [
                "Real Linux auditd ingestion parser correlating SYSCALL, EXECVE, PATH, and CWD records into unified event objects.",
                "Fitted baseline Isolation Forest on 17,846 real events (70% of 25,495 Ubuntu dataset) with persistent joblib model bundle.",
                "Standalone inference engine (predict.py) scoring fresh endpoint telemetry with ZERO retraining or refitting.",
                "Real attack script executions captured and evaluated against live VM auditd logs (find, python3, tar, awk GTFOBins)."
            ]
        ),
        (
            "SYNTHETIC-ONLY (BENCHMARK VALIDATION)",
            "Development Architecture Testing",
            BADGE_SYNTH_BG,
            BADGE_SYNTH_BORDER,
            BADGE_SYNTH_TEXT,
            [
                "Three-tier attack complexity matrix (Tier 1 overt, Tier 2 chained, Tier 3 obfuscated) evaluated on synthetic dataset.",
                "Supervised vs. unsupervised generalization experiment reproducing Stamp (2022) with SMOTE comparators.",
                "Sliding-window size ablation study across N=3, 4, 5 on synthetic multi-step sequences.",
                "Zero-shot LLM attribution tested across synthetic scenario payloads."
            ]
        ),
        (
            "PHASE 2 ROADMAP (DISSERTATION TA2)",
            "Future Work & Engineering Upgrades",
            BADGE_NEUTRAL_BG,
            BADGE_NEUTRAL_BORDER,
            BADGE_NEUTRAL_TEXT,
            [
                "LOLBin argument semantic embeddings & AST-level parsing to eliminate blind spots on subtle syntax (e.g., awk).",
                "Sequence transition-gradient features (binary_changed, entropy_delta) to detect low-signal reconnaissance chains.",
                "Real-data capture for multi-step chained and obfuscated attacks executed directly inside isolated lab VMs.",
                "Cross-OS Windows Sysmon adapter integration (EventID 1 process creation, EventID 11 file creation)."
            ]
        )
    ]

    for i, (title, subtitle, bg, border, text_c, items) in enumerate(status_cols):
        left = Inches(0.8 + i * 3.98)
        card = slide5.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.85), Inches(3.78), Inches(4.85)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = border
        card.line.width = Pt(1.5)

        # Header box
        hdr = slide5.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.12), Inches(1.97), Inches(3.54), Inches(0.65)
        )
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = bg
        hdr.line.color.rgb = border
        htf = hdr.text_frame
        htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_th = htf.paragraphs[0]
        p_th.text = title
        p_th.font.name = FONT_HEADING
        p_th.font.size = Pt(9.5)
        p_th.font.bold = True
        p_th.font.color.rgb = text_c
        p_th.alignment = PP_ALIGN.CENTER

        p_ts = htf.add_paragraph()
        p_ts.text = subtitle
        p_ts.font.name = FONT_BODY
        p_ts.font.size = Pt(8)
        p_ts.font.color.rgb = text_c
        p_ts.alignment = PP_ALIGN.CENTER

        # List items
        tx_list = slide5.shapes.add_textbox(left + Inches(0.15), Inches(2.72), Inches(3.48), Inches(3.85))
        ltf = tx_list.text_frame
        ltf.word_wrap = True

        for j, item in enumerate(items):
            p_item = ltf.paragraphs[0] if j == 0 else ltf.add_paragraph()
            p_item.alignment = PP_ALIGN.LEFT
            p_item.text = f"•  {item}"
            p_item.font.name = FONT_BODY
            p_item.font.size = Pt(9)
            p_item.font.color.rgb = COLOR_SLATE_BODY
            p_item.space_after = Pt(8)

    add_page_decorations(slide5, 5)

    # ==========================================================================
    # SLIDE 6: REAL DATA RESULT #1: BASELINE VALIDATION
    # ==========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(
        slide6,
        title="Real Data Result: Baseline Validation & FPR Calibration",
        subtitle="Unsupervised Isolation Forest evaluation on 25,495 captured Linux auditd telemetry events",
        badge_type="real",
        badge_text="● REAL TELEMETRY: LINUX AUDITD (UBUNTU BASELINE)"
    )

    # Left Column: Metrics & Analysis Card
    left_w = Inches(5.6)
    card6 = slide6.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), left_w, Inches(4.85)
    )
    card6.fill.solid()
    card6.fill.fore_color.rgb = COLOR_BG_CARD
    card6.line.color.rgb = COLOR_BORDER_CARD
    card6.line.width = Pt(1)

    c6_tf = card6.text_frame
    c6_tf.vertical_anchor = MSO_ANCHOR.TOP
    c6_tf.margin_left = Inches(0.25)
    c6_tf.margin_top = Inches(0.2)
    c6_tf.word_wrap = True

    p_6_h = c6_tf.paragraphs[0]
    p_6_h.alignment = PP_ALIGN.LEFT
    p_6_h.text = "Held-Out Baseline Score Distribution (N = 7,649)"
    p_6_h.font.name = FONT_HEADING
    p_6_h.font.size = Pt(11.5)
    p_6_h.font.bold = True
    p_6_h.font.color.rgb = COLOR_NAVY_DARK

    b_stats = [
        ("Total Telemetry Captured:", "25,495 correlated auditd events (Ubuntu)"),
        ("Experimental Split:", "70% Train (17,846 ev) / 30% Held-Out Test (7,649 ev)"),
        ("Empirical FPR on Held-Out:", "4.67% (357 false positives / 7,649 events)"),
        ("Median Score (50th %ile):", "0.0337 (Deep interior of benign hyperspace)"),
        ("90th / 95th Percentiles:", "0.2618 / 0.4564 (Safely below 0.50 threshold)"),
        ("Strongly Benign Bins (<0.2):", "87.36% (6,682 events) score below 0.20"),
        ("Sub-Threshold Normal (<0.5):", "95.34% of benign telemetry falls below cutoff"),
        ("Sliding Window Advantage:", "FPR drops from 7.57% (isolated) to 4.67% (window)")
    ]

    for label, val in b_stats:
        p_row = c6_tf.add_paragraph()
        p_row.alignment = PP_ALIGN.LEFT
        p_row.text = f"•  {label} "
        p_row.font.name = FONT_BODY
        p_row.font.size = Pt(9.5)
        p_row.font.bold = True
        p_row.font.color.rgb = COLOR_NAVY_MID
        p_row.space_before = Pt(4)

        run = p_row.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = COLOR_SLATE_BODY

    # Add Per-Binary FPR note
    p_bin = c6_tf.add_paragraph()
    p_bin.alignment = PP_ALIGN.LEFT
    p_bin.text = "\nPer-Binary Baseline False Positive Discrimination:\n• find: 0.0% FPR (0/66)   |  awk: 0.0% FPR (0/381)\n• tar: 4.17% FPR (1/24)   |  python3: 25.39% FPR (49/193)"
    p_bin.font.name = FONT_BODY
    p_bin.font.size = Pt(8.5)
    p_bin.font.bold = True
    p_bin.font.color.rgb = COLOR_SLATE_MUTED

    # Right Column: Card framing the Chart
    card6_r = slide6.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.85), Inches(5.933), Inches(4.85)
    )
    card6_r.fill.solid()
    card6_r.fill.fore_color.rgb = COLOR_BG_CARD
    card6_r.line.color.rgb = COLOR_BORDER_CARD
    card6_r.line.width = Pt(1)

    chart6_path = "presentation/assets/slide6_baseline_dist.png"
    if os.path.exists(chart6_path):
        slide6.shapes.add_picture(chart6_path, Inches(6.7), Inches(2.0), width=Inches(5.733))

    add_page_decorations(slide6, 6)

    # ==========================================================================
    # SLIDE 7: REAL DATA RESULT #2: LIVE GTFOBIN ATTACK DETECTION
    # ==========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(
        slide7,
        title="Real Data Result: Live GTFOBin Detection & Context Dynamics",
        subtitle="Standalone inference (predict.py) flags real GTFOBins and captures sliding-window elevation",
        badge_type="real",
        badge_text="● REAL TELEMETRY: LIVE GTFOBIN EXECUTION AUDITD"
    )

    # Left Column: Findings Card
    card7 = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(5.6), Inches(4.85)
    )
    card7.fill.solid()
    card7.fill.fore_color.rgb = COLOR_BG_CARD
    card7.line.color.rgb = COLOR_BORDER_CARD
    card7.line.width = Pt(1)

    c7_tf = card7.text_frame
    c7_tf.vertical_anchor = MSO_ANCHOR.TOP
    c7_tf.margin_left = Inches(0.25)
    c7_tf.margin_top = Inches(0.2)
    c7_tf.word_wrap = True

    p_7_h = c7_tf.paragraphs[0]
    p_7_h.alignment = PP_ALIGN.LEFT
    p_7_h.text = "Inference on Live Attack Log (data/live-attack-benign.log)"
    p_7_h.font.name = FONT_HEADING
    p_7_h.font.size = Pt(11.5)
    p_7_h.font.bold = True
    p_7_h.font.color.rgb = COLOR_NAVY_DARK

    findings7 = [
        ("Overt GTFOBin Invocations Successfully Detected:", [
            "sudo find . -exec /bin/sh...  ➔ Anomaly Score: 0.7424 – 0.8503 (FLAGGED)",
            "sudo python3 -c import os... ➔ Anomaly Score: 0.7007 – 0.8132 (FLAGGED)",
            "tar --checkpoint-action=exec... ➔ Anomaly Score: 0.8036 – 0.9702 (FLAGGED)"
        ]),
        ("The 'id' Command Context Experiment (Springer 2026 Validation):", [
            "Normal / Isolated Context: 'id' executed during benign baseline scores 0.0509 – 0.1287 (Events #68, #80, #85, #109 — Sub-threshold Normal).",
            "Mid-Attack Sequence: Exactly the same 'id' command executed immediately following tar checkpoint or hex shell jumps to 0.5792 (#83) and 0.5834 (#106) — CROSSING threshold!",
            "Concrete Proof: The sliding window correctly captures multi-step behavioral elevation rather than isolated command text."
        ])
    ]

    for title, subitems in findings7:
        p_th = c7_tf.add_paragraph()
        p_th.alignment = PP_ALIGN.LEFT
        p_th.text = f"\n{title}"
        p_th.font.name = FONT_BODY
        p_th.font.size = Pt(9.5)
        p_th.font.bold = True
        p_th.font.color.rgb = COLOR_NAVY_MID

        for sub in subitems:
            p_sub = c7_tf.add_paragraph()
            p_sub.alignment = PP_ALIGN.LEFT
            p_sub.text = f"• {sub}"
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(8.5)
            p_sub.font.color.rgb = COLOR_SLATE_BODY
            p_sub.space_before = Pt(2.5)

    # Right Column: Card framing the Chart
    card7_r = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.85), Inches(5.933), Inches(4.85)
    )
    card7_r.fill.solid()
    card7_r.fill.fore_color.rgb = COLOR_BG_CARD
    card7_r.line.color.rgb = COLOR_BORDER_CARD
    card7_r.line.width = Pt(1)

    chart7_path = "presentation/assets/slide7_window_context.png"
    if os.path.exists(chart7_path):
        slide7.shapes.add_picture(chart7_path, Inches(6.7), Inches(2.0), width=Inches(5.733))

    add_page_decorations(slide7, 7)

    # ==========================================================================
    # SLIDE 8: HONEST LIMITATION #1: UNDETECTED AWK
    # ==========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(
        slide8,
        title="Honest Limitation 1: Undetected GTFOBin (awk)",
        subtitle="The awk-based privilege escalation technique scored strictly below 0.50 and evaded detection",
        badge_type="limit",
        badge_text="■ HONEST SCIENTIFIC LIMITATION (REAL AUDITD FINDING)"
    )

    card8_top = slide8.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.733), Inches(1.35)
    )
    card8_top.fill.solid()
    card8_top.fill.fore_color.rgb = COLOR_BG_CARD
    card8_top.line.color.rgb = BADGE_LIMIT_BORDER
    card8_top.line.width = Pt(1.5)

    ctf8 = card8_top.text_frame
    ctf8.vertical_anchor = MSO_ANCHOR.TOP
    ctf8.margin_left = Inches(0.2)
    ctf8.margin_top = Inches(0.12)

    p_8t = ctf8.paragraphs[0]
    p_8t.alignment = PP_ALIGN.LEFT
    p_8t.text = "Concrete Empirical Finding: sudo awk BEGIN {system(...)} Was Not Detected"
    p_8t.font.name = FONT_HEADING
    p_8t.font.size = Pt(12)
    p_8t.font.bold = True
    p_8t.font.color.rgb = BADGE_LIMIT_TEXT

    p_8desc = ctf8.add_paragraph()
    p_8desc.alignment = PP_ALIGN.LEFT
    p_8desc.text = "During live attack simulation (live-attack-benign.log, Events #57–#66), the awk-based shell spawn technique produced scores between 0.0617 and 0.4928 across all attempts. Every single invocation remained sub-threshold (< 0.50 decision boundary), representing a complete false negative."
    p_8desc.font.name = FONT_BODY
    p_8desc.font.size = Pt(9.5)
    p_8desc.font.color.rgb = COLOR_SLATE_BODY
    p_8desc.space_before = Pt(3)

    # Three explanatory columns
    reasons = [
        (
            "1. Benign Distribution Masking",
            "• awk appears 381 times in the held-out benign baseline.\n• The model established a strong prior for awk as routine administrative text processing.\n• Table 2C shows awk has 0.0000 FPR on benign data; the detector heavily normalizes awk activity."
        ),
        (
            "2. Feature Space Insensitivity",
            "• Syntax 'BEGIN {system(...)}' does not contain explicit exploit tokens like '-exec', 'checkpoint', or 'bypass'.\n• Character Shannon entropy remains moderate (~3.2–3.6), indistinguishable from complex sed/awk queries.\n• Token hit counter recorded 0 hits."
        ),
        (
            "3. Privilege Signal Latency",
            "• auditd captured the awk process execution under sudo, but the child shell was spawned internally.\n• The instantaneous command did not trigger a logged UID change at the primary audit event boundary.\n• Generic lineage features failed to detect execution pivot."
        )
    ]

    for i, (rtitle, rdesc) in enumerate(reasons):
        left = Inches(0.8 + i * 3.98)
        rcard = slide8.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(3.40), Inches(3.78), Inches(2.45)
        )
        rcard.fill.solid()
        rcard.fill.fore_color.rgb = COLOR_BG_CARD
        rcard.line.color.rgb = COLOR_BORDER_CARD
        rcard.line.width = Pt(1)

        rtf = rcard.text_frame
        rtf.vertical_anchor = MSO_ANCHOR.TOP
        rtf.margin_left = Inches(0.18)
        rtf.margin_top = Inches(0.14)

        p_rh = rtf.paragraphs[0]
        p_rh.alignment = PP_ALIGN.LEFT
        p_rh.text = rtitle
        p_rh.font.name = FONT_HEADING
        p_rh.font.size = Pt(10.5)
        p_rh.font.bold = True
        p_rh.font.color.rgb = COLOR_NAVY_DARK

        p_rb = rtf.add_paragraph()
        p_rb.alignment = PP_ALIGN.LEFT
        p_rb.text = rdesc
        p_rb.font.name = FONT_BODY
        p_rb.font.size = Pt(9)
        p_rb.font.color.rgb = COLOR_SLATE_BODY
        p_rb.space_before = Pt(5)

    # Takeaway banner
    takeaway_card = slide8.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.65)
    )
    takeaway_card.fill.solid()
    takeaway_card.fill.fore_color.rgb = RGBColor(241, 245, 249)
    takeaway_card.line.color.rgb = COLOR_BORDER_CARD
    ttf = takeaway_card.text_frame
    ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ttf.margin_left = Inches(0.2)
    p_tkw = ttf.paragraphs[0]
    p_tkw.text = "Key Academic Takeaway: Generic behavioral features struggle with syntax-subtle LOLBins. Catching awk requires semantic argument modeling and token n-gram embeddings (Phase 2), not just generic command length and entropy."
    p_tkw.font.name = FONT_BODY
    p_tkw.font.size = Pt(9.5)
    p_tkw.font.bold = True
    p_tkw.font.color.rgb = COLOR_NAVY_MID

    add_page_decorations(slide8, 8)

    # ==========================================================================
    # SLIDE 9: HONEST LIMITATION #2: LOW-SIGNAL RECONNAISSANCE CHAIN
    # ==========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(
        slide9,
        title="Honest Limitation 2: Low-Signal Reconnaissance Blind Spot",
        subtitle="Multi-step discovery chains evade detection when every command in the chain is individually low-signal",
        badge_type="limit",
        badge_text="■ HONEST SCIENTIFIC LIMITATION (SEQUENCE FEATURE GAP)"
    )

    card9_top = slide9.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(11.733), Inches(1.35)
    )
    card9_top.fill.solid()
    card9_top.fill.fore_color.rgb = COLOR_BG_CARD
    card9_top.line.color.rgb = BADGE_LIMIT_BORDER
    card9_top.line.width = Pt(1.5)

    ctf9 = card9_top.text_frame
    ctf9.vertical_anchor = MSO_ANCHOR.TOP
    ctf9.margin_left = Inches(0.2)
    ctf9.margin_top = Inches(0.12)

    p_9t = ctf9.paragraphs[0]
    p_9t.alignment = PP_ALIGN.LEFT
    p_9t.text = "Concrete Empirical Finding: whoami ➔ id ➔ hostname ➔ uname -a Chain Was Missed"
    p_9t.font.name = FONT_HEADING
    p_9t.font.size = Pt(12)
    p_9t.font.bold = True
    p_9t.font.color.rgb = BADGE_LIMIT_TEXT

    p_9desc = ctf9.add_paragraph()
    p_9desc.alignment = PP_ALIGN.LEFT
    p_9desc.text = "In both baseline logs and live simulation, a sequence composed purely of standard discovery utilities (whoami, id, hostname, uname -a, groups, ls, pwd) produced anomaly scores strictly between 0.01 and 0.13. The sliding window failed to trigger an alert despite coherent reconnaissance activity."
    p_9desc.font.name = FONT_BODY
    p_9desc.font.size = Pt(9.5)
    p_9desc.font.color.rgb = COLOR_SLATE_BODY
    p_9desc.space_before = Pt(3)

    recon_analysis = [
        (
            "Order-Invariant Feature Aggregation",
            "• Current sliding window computes scalar aggregates: max, avg, and sum.\n• It lacks transition dynamics (e.g., priv_transition_delta, binary_changed, entropy_gradient).\n• The window treats [whoami, id, uname] identically to repeated benign commands."
        ),
        (
            "Dependence on an Overt 'Anchor' Step",
            "• Synthetic Tier 2 chains reached 100% chain interception ONLY because each chain contained at least one overt high-signal step (e.g., sudo tar checkpoint).\n• The overt step triggered the window; the window did not detect the subtle reconnaissance steps."
        ),
        (
            "Lack of Sequence Transition Modeling",
            "• Isolation Forest partitions vector coordinates; it does not model Markov state transitions or temporal progression.\n• When all steps have 0 token hits and no privilege change, window cumulative sums remain safely below the 0.50 threshold."
        )
    ]

    for i, (rtitle, rdesc) in enumerate(recon_analysis):
        left = Inches(0.8 + i * 3.98)
        rcard = slide9.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(3.40), Inches(3.78), Inches(2.45)
        )
        rcard.fill.solid()
        rcard.fill.fore_color.rgb = COLOR_BG_CARD
        rcard.line.color.rgb = COLOR_BORDER_CARD
        rcard.line.width = Pt(1)

        rtf = rcard.text_frame
        rtf.vertical_anchor = MSO_ANCHOR.TOP
        rtf.margin_left = Inches(0.18)
        rtf.margin_top = Inches(0.14)

        p_rh = rtf.paragraphs[0]
        p_rh.alignment = PP_ALIGN.LEFT
        p_rh.text = rtitle
        p_rh.font.name = FONT_HEADING
        p_rh.font.size = Pt(10.5)
        p_rh.font.bold = True
        p_rh.font.color.rgb = COLOR_NAVY_DARK

        p_rb = rtf.add_paragraph()
        p_rb.alignment = PP_ALIGN.LEFT
        p_rb.text = rdesc
        p_rb.font.name = FONT_BODY
        p_rb.font.size = Pt(9)
        p_rb.font.color.rgb = COLOR_SLATE_BODY
        p_rb.space_before = Pt(5)

    takeaway_card9 = slide9.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.65)
    )
    takeaway_card9.fill.solid()
    takeaway_card9.fill.fore_color.rgb = RGBColor(241, 245, 249)
    takeaway_card9.line.color.rgb = COLOR_BORDER_CARD
    ttf9 = takeaway_card9.text_frame
    ttf9.vertical_anchor = MSO_ANCHOR.MIDDLE
    ttf9.margin_left = Inches(0.2)
    p_tkw9 = ttf9.paragraphs[0]
    p_tkw9.text = "Key Academic Takeaway: Unsupervised Isolation Forest cannot infer malicious recon intent from zero-signal commands without explicit sequence transition features (binary_changed, token_acceleration), which form our primary Phase 2 objective."
    p_tkw9.font.name = FONT_BODY
    p_tkw9.font.size = Pt(9.5)
    p_tkw9.font.bold = True
    p_tkw9.font.color.rgb = COLOR_NAVY_MID

    add_page_decorations(slide9, 9)

    # ==========================================================================
    # SLIDE 10: SYNTHETIC ARCHITECTURE VALIDATION (SECONDARY, BRIEF)
    # ==========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(
        slide10,
        title="Synthetic Validation: Multi-Tier Sensitivity & Overfitting",
        subtitle="Controlled architectural experiments conducted during development to validate theoretical models",
        badge_type="synth",
        badge_text="▲ SYNTHETIC DATA: ARCHITECTURE VALIDATION BENCHMARK"
    )

    # Left Column: Structured Results Card
    card10 = slide10.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(5.6), Inches(4.85)
    )
    card10.fill.solid()
    card10.fill.fore_color.rgb = COLOR_BG_CARD
    card10.line.color.rgb = COLOR_BORDER_CARD
    card10.line.width = Pt(1)

    c10_tf = card10.text_frame
    c10_tf.vertical_anchor = MSO_ANCHOR.TOP
    c10_tf.margin_left = Inches(0.25)
    c10_tf.margin_top = Inches(0.2)
    c10_tf.word_wrap = True

    p_10_h = c10_tf.paragraphs[0]
    p_10_h.alignment = PP_ALIGN.LEFT
    p_10_h.text = "Controlled Lab Experiments (Synthetic Sysmon + Auditd)"
    p_10_h.font.name = FONT_HEADING
    p_10_h.font.size = Pt(11.5)
    p_10_h.font.bold = True
    p_10_h.font.color.rgb = COLOR_NAVY_DARK

    synth_findings = [
        ("Three Attack Complexity Tiers (SLR 2026 Validation):", [
            "Tier 1 (Overt / Isolated): Recall 100.0%, Precision 0.9901, F1 0.9950 (300 events)",
            "Tier 2 (Chained Multi-Step): Recall 91.23%, Precision 0.9738, F1 0.9420 (285 events)",
            "Tier 3 (Obfuscated / Stealthy): Recall 90.00%, Precision 0.9507, F1 0.9247 (150 events)",
            "Empirical confirmation: Monotonic degradation matches SLR 2026 theorem on obfuscation limits."
        ]),
        ("Supervised Generalization Failure (Reproducing Stamp 2022):", [
            "Isolation Forest (Unsupervised): 90.80% Recall on unseen Tiers 2 & 3 with ZERO attack labels in training.",
            "Random Forest (Supervised + SMOTE): Achieves 100% on seen Tier 1, but collapses to 48.28% Recall on unseen Tiers 2 & 3.",
            "Decision Tree (Supervised + SMOTE): Collapses to 51.72% Recall on unseen Tiers 2 & 3.",
            "Scientific Validation: Proves supervised classifiers heavily overfit to seen syntax, validating our unsupervised design."
        ])
    ]

    for title, subitems in synth_findings:
        p_th = c10_tf.add_paragraph()
        p_th.alignment = PP_ALIGN.LEFT
        p_th.text = f"\n{title}"
        p_th.font.name = FONT_BODY
        p_th.font.size = Pt(9.5)
        p_th.font.bold = True
        p_th.font.color.rgb = COLOR_NAVY_MID

        for sub in subitems:
            p_sub = c10_tf.add_paragraph()
            p_sub.alignment = PP_ALIGN.LEFT
            p_sub.text = f"• {sub}"
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(8.5)
            p_sub.font.color.rgb = COLOR_SLATE_BODY
            p_sub.space_before = Pt(2.5)

    # Right Column: Card framing the Chart
    card10_r = slide10.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.85), Inches(5.933), Inches(4.85)
    )
    card10_r.fill.solid()
    card10_r.fill.fore_color.rgb = COLOR_BG_CARD
    card10_r.line.color.rgb = COLOR_BORDER_CARD
    card10_r.line.width = Pt(1)

    chart10_path = "presentation/assets/slide10_synthetic_eval.png"
    if os.path.exists(chart10_path):
        slide10.shapes.add_picture(chart10_path, Inches(6.7), Inches(2.0), width=Inches(5.733))

    add_page_decorations(slide10, 10)

    # ==========================================================================
    # SLIDE 11: TOOLS & TECHNOLOGY
    # ==========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(
        slide11,
        title="Tools, Technology & Lab Infrastructure",
        subtitle="Modular, reproducible engineering stack utilizing standard open-source systems and ML frameworks",
        badge_type="neutral",
        badge_text="◆ ENGINEERING STACK & LAB ENVIRONMENT"
    )

    stack_cards = [
        (
            "Lab Virtualization & Telemetry",
            "Isolated VM Simulation Environment",
            [
                "Hypervisor: KVM / QEMU managed via virt-manager on Linux.",
                "Target VM: Ubuntu Server 24.04 LTS (isolated, non-networked).",
                "Telemetry Sensor: Linux auditd capturing SYSCALL, EXECVE, PATH, CWD.",
                "Safe Artifacts: bash and powershell simulation scripts decoupled from host."
            ]
        ),
        (
            "Machine Learning & Pipeline Core",
            "Unsupervised Anomaly Engine",
            [
                "Python 3.14 / 3.12 scientific compute environment.",
                "scikit-learn: IsolationForest (150 estimators), StandardScaler.",
                "imbalanced-learn: SMOTE used strictly for supervised benchmark comparisons.",
                "joblib: Standalone pipeline bundle serialization (detector_model.joblib)."
            ]
        ),
        (
            "Feature Engineering & Windowing",
            "Sequence Behavioral Analytics",
            [
                "Custom continuous Shannon entropy & character dynamics calculation.",
                "Sliding Command Window (N=3 to 5) aggregating session context.",
                "Process lineage rarity mapping (parent-child co-occurrence frequency).",
                "Dual-OS token dictionaries (Linux GTFOBins & Windows LOLBins)."
            ]
        ),
        (
            "Explainability & SOC Console",
            "Zero-Shot Cognitive Layer",
            [
                "LLM API Integration: OpenAI / Anthropic / Groq REST endpoints.",
                "Strict Zero-Shot Prompt: Constrained to JSON feature anomalies (no hallucination).",
                "Deterministic Offline SOC Fallback: Full testability without API keys.",
                "Streamlit Dashboard: Interactive triage console with telemetry toggle."
            ]
        )
    ]

    for i, (title, subtitle, items) in enumerate(stack_cards):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.85 + row * 2.45)

        card = slide11.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(2.3)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_BORDER_CARD
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.margin_left = Inches(0.2)
        ctf.margin_top = Inches(0.12)

        p_t = ctf.paragraphs[0]
        p_t.alignment = PP_ALIGN.LEFT
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(11.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY_DARK

        p_s = ctf.add_paragraph()
        p_s.alignment = PP_ALIGN.LEFT
        p_s.text = subtitle
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(9.5)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_ACCENT_BLUE
        p_s.space_after = Pt(4)

        for item in items:
            p_i = ctf.add_paragraph()
            p_i.alignment = PP_ALIGN.LEFT
            p_i.text = f"•  {item}"
            p_i.font.name = FONT_BODY
            p_i.font.size = Pt(9)
            p_i.font.color.rgb = COLOR_SLATE_BODY
            p_i.space_before = Pt(2)

    add_page_decorations(slide11, 11)

    # ==========================================================================
    # SLIDE 12: NEXT STEPS (PHASE 2 ROADMAP)
    # ==========================================================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(
        slide12,
        title="Research Roadmap: Transition to Phase 2 (TA2)",
        subtitle="Evolving from generic behavioral anomaly detection to LOLBin-specific syntactic modeling",
        badge_type="neutral",
        badge_text="◆ DISSERTATION ROADMAP (SEMESTER IV PLAN)"
    )

    roadmap_steps = [
        (
            "1. LOLBin Argument Semantic Modeling",
            "Objective: Eliminate blind spots like awk",
            "Transition from generic command entropy to semantic token embeddings (Ongun et al. 2021) and AST argument parsing. Distinguish routine text parsing from subshell invocations (e.g., BEGIN {system()})."
        ),
        (
            "2. Sequence Transition-Gradient Features",
            "Objective: Detect low-signal discovery chains",
            "Augment the sliding window with dynamic inter-command features: binary_changed flag, token_acceleration, and session entropy gradients to catch reconnaissance chains before privilege escalation occurs."
        ),
        (
            "3. Multi-Tier Real VM Telemetry Capture",
            "Objective: Replace synthetic benchmark with real VM traces",
            "Execute full Tier 2 (chained sequences) and Tier 3 (obfuscated GTFOBins) attack matrices inside isolated Ubuntu VMs, capturing 100% real auditd logs to validate production generalization."
        ),
        (
            "4. Cross-OS Windows Sysmon Integration",
            "Objective: Dual-OS EDR coverage",
            "Implement a dedicated Windows Sysmon XML/JSON parser (EventID 1 process creation, EventID 11 file creation) and benchmark prevalent Windows LOLBins (certutil, powershell, mshta, rundll32)."
        )
    ]

    for i, (title, obj, desc) in enumerate(roadmap_steps):
        top_pos = Inches(1.85 + i * 1.22)
        card = slide12.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.733), Inches(1.08)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = COLOR_BORDER_CARD
        card.line.width = Pt(1)

        # Left indicator tag
        tag = slide12.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), top_pos, Inches(0.1), Inches(1.08)
        )
        tag.fill.solid()
        tag.fill.fore_color.rgb = COLOR_NAVY_MID
        tag.line.fill.background()

        ctf = card.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.12)

        p_t = ctf.paragraphs[0]
        p_t.alignment = PP_ALIGN.LEFT
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(11.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_NAVY_DARK

        p_o = ctf.add_paragraph()
        p_o.alignment = PP_ALIGN.LEFT
        p_o.text = f"{obj} — {desc}"
        p_o.font.name = FONT_BODY
        p_o.font.size = Pt(9.5)
        p_o.font.color.rgb = COLOR_SLATE_BODY
        p_o.space_before = Pt(3)

    add_page_decorations(slide12, 12)

    # Save presentation
    os.makedirs(os.path.dirname(output_pptx), exist_ok=True)
    prs.save(output_pptx)
    print(f"[+] Successfully generated PowerPoint presentation: {output_pptx}")


if __name__ == "__main__":
    create_deck()
